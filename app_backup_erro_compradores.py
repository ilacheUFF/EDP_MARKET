import streamlit as st
import pandas as pd
from pathlib import Path
from openpyxl import load_workbook
from datetime import datetime

ARQUIVO = Path("EDP_MARKET.xlsx")

st.set_page_config(
    page_title="EDP Market",
    page_icon="🌱",
    layout="wide"
)


# ============================================================
# FUNÇÕES DE DADOS
# ============================================================

@st.cache_data

def carregar_dados():
    if not ARQUIVO.exists():
        return None, None

    alunos = pd.read_excel(ARQUIVO, sheet_name="Alunos")
    edps = pd.read_excel(ARQUIVO, sheet_name="EDPs")

    alunos["EDP atribuído"] = (
        alunos["EDP atribuído"].fillna("").astype(str).str.strip().str.upper()
    )
    edps["Código"] = (
        edps["Código"].fillna("").astype(str).str.strip().str.upper()
    )

    # Garante as colunas necessárias, caso a planilha ainda esteja na versão inicial.
    if "Saldo" not in alunos.columns:
        alunos["Saldo"] = alunos.get("Orçamento inicial", 1000)
    if "Pontos" not in alunos.columns:
        alunos["Pontos"] = 0

    if "Vendedor" not in edps.columns:
        edps["Vendedor"] = ""
    if "Preço" not in edps.columns:
        edps["Preço"] = None
    if "Status" not in edps.columns:
        edps["Status"] = "Não anunciado"

    return alunos, edps


def salvar_aba(df, nome_aba):
    """Atualiza apenas uma aba do Excel, preservando as demais."""
    if not ARQUIVO.exists():
        st.error("Arquivo EDP_MARKET.xlsx não encontrado.")
        return False

    try:
        # openpyxl permite alterar uma aba sem apagar as outras.
        wb = load_workbook(ARQUIVO)

        if nome_aba in wb.sheetnames:
            ws = wb[nome_aba]
            wb.remove(ws)

        ws = wb.create_sheet(nome_aba)

        # Cabeçalho
        for col_idx, coluna in enumerate(df.columns, start=1):
            ws.cell(row=1, column=col_idx, value=coluna)

        # Dados
        for row_idx, (_, linha) in enumerate(df.iterrows(), start=2):
            for col_idx, valor in enumerate(linha, start=1):
                # NaN/NaT não devem ir para o Excel.
                if pd.isna(valor):
                    valor = None
                ws.cell(row=row_idx, column=col_idx, value=valor)

        wb.save(ARQUIVO)
        return True

    except PermissionError:
        st.error(
            "❌ Não foi possível salvar o Excel. "
            "Verifique se o arquivo EDP_MARKET.xlsx está aberto em outro programa."
        )
        return False
    except Exception as e:
        st.error(f"❌ Erro ao salvar o Excel: {e}")
        return False


def atualizar_edp_no_excel(codigo, vendedor, preco):
    """Coloca um EDP à venda diretamente pelo Streamlit."""
    try:
        wb = load_workbook(ARQUIVO)
        ws = wb["EDPs"]

        # Localiza as colunas pelo nome.
        cabecalho = {
            str(ws.cell(row=1, column=c).value).strip(): c
            for c in range(1, ws.max_column + 1)
        }

        col_codigo = cabecalho["Código"]
        col_vendedor = cabecalho["Vendedor"]
        col_preco = cabecalho["Preço"]
        col_status = cabecalho["Status"]

        for linha in range(2, ws.max_row + 1):
            valor_codigo = str(ws.cell(row=linha, column=col_codigo).value or "").strip().upper()

            if valor_codigo == codigo:
                ws.cell(row=linha, column=col_vendedor, value=vendedor)
                ws.cell(row=linha, column=col_preco, value=float(preco))
                ws.cell(row=linha, column=col_status, value="À venda")
                wb.save(ARQUIVO)
                return True

        return False

    except PermissionError:
        st.error(
            "❌ Não foi possível salvar o preço. "
            "Feche o EDP_MARKET.xlsx caso ele esteja aberto no Excel."
        )
        return False
    except Exception as e:
        st.error(f"❌ Erro ao atualizar o EDP: {e}")
        return False


# ============================================================
# CARREGAMENTO
# ============================================================

alunos, edps = carregar_dados()

if alunos is None or edps is None:
    st.error("Não foi possível encontrar a planilha EDP_MARKET.xlsx.")
    st.stop()

edp_lookup = (
    edps.drop_duplicates("Código")
    .set_index("Código")[["EDP", "Etapa"]]
    .to_dict("index")
)

alunos["EDP"] = alunos["EDP atribuído"].map(
    lambda codigo: edp_lookup.get(codigo, {}).get("EDP", "") if codigo else ""
)
alunos["Etapa automática"] = alunos["EDP atribuído"].map(
    lambda codigo: edp_lookup.get(codigo, {}).get("Etapa", "") if codigo else ""
)

menu = st.sidebar.radio(
    "Navegação",
    ["🎮 Início", "👥 Alunos", "🌱 EDPs", "🛒 Mercado"]
)


# ============================================================
# INÍCIO
# ============================================================

if menu == "🎮 Início":

    st.title("🌱 EDP MARKET")
    st.subheader("Mercado de Decisões em Ecodesign")

    col1, col2, col3, col4 = st.columns(4)

    with col1:
        st.metric("EDPs", len(edps))
    with col2:
        st.metric("Alunos", len(alunos))
    with col3:
        st.metric("Etapas", 5)
    with col4:
        st.metric("Orçamento inicial", "1000 EC")

    st.divider()

    atribuidos = int((alunos["EDP atribuído"] != "").sum())
    validos = int(
        ((alunos["EDP atribuído"] != "") &
         (alunos["EDP"].fillna("") != "")).sum()
    )

    c1, c2 = st.columns(2)

    with c1:
        st.metric("Alunos com EDP atribuído", f"{atribuidos}/{len(alunos)}")
    with c2:
        st.metric("Atribuições válidas", f"{validos}/{len(alunos)}")

    if validos == len(alunos):
        st.success("✅ Todos os alunos estão associados a um EDP válido.")
    else:
        st.warning(
            "⚠️ Existem alunos sem EDP ou com código de EDP não encontrado."
        )

    st.info(
        "Durante a aula, o professor poderá definir o preço de cada EDP "
        "diretamente no Streamlit. O Excel será usado como banco de dados, "
        "sem necessidade de ficar alternando de tela."
    )


# ============================================================
# ALUNOS
# ============================================================

elif menu == "👥 Alunos":

    st.title("👥 Alunos")
    st.write(f"**Total cadastrado:** {len(alunos)} alunos")

    colunas = [
        "ID", "Nome", "EDP atribuído", "EDP",
        "Etapa automática", "Saldo", "Pontos"
    ]

    st.dataframe(
        alunos[colunas],
        use_container_width=True,
        hide_index=True
    )

    st.divider()
    st.subheader("🔎 Consultar aluno")

    nomes_alunos = alunos["Nome"].dropna().tolist()

    if nomes_alunos:

        nome = st.selectbox("Selecione um aluno", nomes_alunos)
        aluno = alunos[alunos["Nome"] == nome].iloc[0]

        c1, c2, c3 = st.columns(3)

        with c1:
            st.metric("ID", int(aluno["ID"]))
        with c2:
            st.metric("Saldo", f'{float(aluno["Saldo"]):.0f} EC')
        with c3:
            st.metric("Pontos", f'{float(aluno["Pontos"]):.0f}')

        st.divider()

        codigo = aluno["EDP atribuído"]

        if codigo and codigo in edp_lookup:

            info = edp_lookup[codigo]

            st.success("✅ EDP identificado automaticamente")
            st.markdown(f"### {codigo}")

            d1, d2 = st.columns(2)

            with d1:
                st.write("**EDP:**")
                st.write(info.get("EDP", "Não informado"))

            with d2:
                st.write("**Etapa do ciclo de vida:**")
                st.write(info.get("Etapa", "Não informada"))

        elif codigo:
            st.error(
                f"❌ O código **{codigo}** não foi encontrado na aba EDPs."
            )
        else:
            st.warning("⚠️ Este aluno ainda não possui um EDP atribuído.")


# ============================================================
# EDPs
# ============================================================

elif menu == "🌱 EDPs":

    st.title("🌱 Eco-design Principles")
    st.write(f"**Total de EDPs:** {len(edps)}")

    st.dataframe(
        edps,
        use_container_width=True,
        hide_index=True
    )

    st.divider()
    st.subheader("🔎 Consultar EDP")

    codigos = edps["Código"].dropna().tolist()

    if codigos:

        codigo = st.selectbox("Selecione um EDP", codigos)
        edp = edps[edps["Código"] == codigo].iloc[0]

        st.markdown(f"### {edp['Código']}")

        c1, c2, c3 = st.columns(3)

        with c1:
            st.write("**EDP:**")
            st.write(edp["EDP"] if pd.notna(edp["EDP"]) else "Não cadastrado")

        with c2:
            st.write("**Etapa:**")
            st.write(edp["Etapa"] if pd.notna(edp["Etapa"]) else "Não definida")

        with c3:
            st.write("**Status:**")
            st.write(edp["Status"] if pd.notna(edp["Status"]) else "Não anunciado")

        vendedores = alunos[alunos["EDP atribuído"] == codigo]

        if len(vendedores) > 0:
            st.divider()
            st.write("**Aluno(s) responsável(is):**")
            st.dataframe(
                vendedores[["ID", "Nome"]],
                use_container_width=True,
                hide_index=True
            )
        else:
            st.info("Nenhum aluno foi associado a este EDP.")


# ============================================================
# MERCADO
# ============================================================

elif menu == "🛒 Mercado":

    st.title("🛒 EDP MARKET")
    st.subheader("Mercado de Decisões em Ecodesign")

    st.info(
        "🎤 Nesta tela você controla a apresentação, define o preço e coloca "
        "o EDP à venda. Depois, no mesmo local, os colegas podem comprar. "
        "Não é necessário abrir o Excel durante a aula."
    )

    # --------------------------------------------------------
    # BLOCO 1 — APRESENTAÇÃO E DEFINIÇÃO DO PREÇO
    # --------------------------------------------------------

    st.markdown("## 🎤 Apresentação e venda")

    nomes_alunos = alunos["Nome"].dropna().tolist()

    if not nomes_alunos:
        st.warning("Nenhum aluno cadastrado.")
        st.stop()

    vendedor_nome = st.selectbox(
        "👤 Quem está apresentando?",
        nomes_alunos,
        key="vendedor_apresentacao"
    )

    vendedor = alunos[alunos["Nome"] == vendedor_nome].iloc[0]
    codigo_vendedor = str(vendedor["EDP atribuído"]).strip().upper()

    if not codigo_vendedor or codigo_vendedor not in edp_lookup:
        st.error("❌ Este aluno não possui um EDP válido atribuído.")
    else:
        info_vendedor = edp_lookup[codigo_vendedor]

        c1, c2 = st.columns([3, 1])

        with c1:
            st.markdown(f"### {codigo_vendedor}")
            st.write(f"**EDP:** {info_vendedor['EDP']}")
            st.write(f"**Etapa:** {info_vendedor['Etapa']}")

        edp_atual = edps[edps["Código"] == codigo_vendedor].iloc[0]
        status_atual = str(edp_atual["Status"] or "Não anunciado").strip()

        with c2:
            st.metric("Status", status_atual)

        if status_atual.lower() == "à venda":
            preco_atual = edp_atual["Preço"]
            if pd.notna(preco_atual):
                st.success(
                    f"🟢 {codigo_vendedor} já está à venda por "
                    f"**{float(preco_atual):.0f} EcoCoins**."
                )
            st.caption("Este EDP já foi anunciado e não precisa ser anunciado novamente.")

        else:
            st.write("**Defina o preço após a apresentação:**")

            preco = st.number_input(
                "Preço do EDP (EcoCoins)",
                min_value=1,
                max_value=1000,
                value=250,
                step=10,
                key=f"preco_{codigo_vendedor}"
            )

            if st.button(
                "🟢 COLOCAR EDP À VENDA",
                type="primary",
                key=f"vender_{codigo_vendedor}"
            ):
                sucesso = atualizar_edp_no_excel(
                    codigo_vendedor,
                    vendedor_nome,
                    preco
                )

                if sucesso:
                    st.success(
                        f"✅ {codigo_vendedor} colocado à venda por "
                        f"**{preco:.0f} EcoCoins**."
                    )
                    st.caption(
                        f"Salvo no Excel às {datetime.now().strftime('%H:%M:%S')}."
                    )
                    st.cache_data.clear()
                    st.rerun()

    st.divider()

    # --------------------------------------------------------
    # BLOCO 2 — MERCADO / COMPRAS
    # --------------------------------------------------------

    st.markdown("## 🛒 Registro rápido das compras")

    st.info(
        "💡 Em vez de registrar uma compra por vez, selecione o EDP vendido "
        "e marque todos os alunos que compraram. Depois clique uma única vez "
        "em **REGISTRAR COMPRADORES**."
    )

    disponiveis = edps[
        edps["Status"].fillna("").astype(str).str.strip().str.lower()
        == "à venda"
    ].copy()

    if len(disponiveis) == 0:
        st.warning("Nenhum EDP está à venda neste momento.")
    else:
        opcoes_edp = {}
        for _, linha in disponiveis.iterrows():
            codigo = str(linha["Código"])
            vendedor_edp = str(linha["Vendedor"])
            preco_edp = float(linha["Preço"])
            opcoes_edp[codigo] = (
                f"{codigo} — {linha['EDP']} | {linha['Etapa']} | "
                f"{preco_edp:.0f} EC | {vendedor_edp}"
            )

        codigo_compra = st.selectbox(
            "🌱 Qual EDP está sendo vendido agora?",
            list(opcoes_edp.keys()),
            format_func=lambda codigo: opcoes_edp[codigo],
            key="edp_registro_venda"
        )

        edp_venda = disponiveis[disponiveis["Código"] == codigo_compra].iloc[0]
        vendedor_edp = str(edp_venda["Vendedor"])
        preco_edp = float(edp_venda["Preço"])

        st.markdown(
            f"### {codigo_compra} — {edp_venda['EDP']}"
        )
        c1, c2, c3 = st.columns(3)
        with c1:
            st.write(f"**Etapa:** {edp_venda['Etapa']}")
        with c2:
            st.write(f"**Vendedor:** {vendedor_edp}")
        with c3:
            st.write(f"**Preço:** {preco_edp:.0f} EcoCoins")

        nomes_compradores = [
            nome for nome in nomes_alunos if nome != vendedor_edp
        ]

        compradores_selecionados = st.multiselect(
            "👥 Marque todos os alunos que compraram este EDP:",
            nomes_compradores,
            placeholder="Clique para selecionar os compradores...",
            key=f"compradores_{codigo_compra}"
        )

        if compradores_selecionados:
            st.success(
                f"{len(compradores_selecionados)} comprador(es) selecionado(s)."
            )

        if st.button(
            "🛒 REGISTRAR COMPRADORES",
            type="primary",
            disabled=len(compradores_selecionados) == 0,
            key=f"registrar_{codigo_compra}"
        ):
            if "compras" not in st.session_state:
                st.session_state.compras = []

            compradores_ja_registrados = {
                (c["Comprador ID"], c["Código"])
                for c in st.session_state.compras
            }

            registrados = 0
            ignorados = 0

            for nome_comprador in compradores_selecionados:
                linha_comprador = alunos[
                    alunos["Nome"] == nome_comprador
                ].iloc[0]
                comprador_id = int(linha_comprador["ID"])
                saldo_atual = float(linha_comprador["Saldo"])
                chave = (comprador_id, codigo_compra)

                if chave in compradores_ja_registrados:
                    ignorados += 1
                    continue

                if preco_edp > saldo_atual:
                    ignorados += 1
                    continue

                st.session_state.compras.append({
                    "Comprador ID": comprador_id,
                    "Comprador": nome_comprador,
                    "Vendedor": vendedor_edp,
                    "Código": codigo_compra,
                    "EDP": str(edp_venda["EDP"]),
                    "Etapa": str(edp_venda["Etapa"]),
                    "Preço": preco_edp
                })
                registrados += 1

            if registrados > 0:
                st.success(
                    f"✅ {registrados} compra(s) registrada(s) para {codigo_compra}."
                )
            if ignorados > 0:
                st.warning(
                    f"⚠️ {ignorados} seleção(ões) não foram registradas "
                    "(compra já registrada ou saldo insuficiente)."
                )

            st.rerun()

    # --------------------------------------------------------
    # COMPRAS DA SESSÃO
    # --------------------------------------------------------

    if "compras" in st.session_state:

        compras = pd.DataFrame(st.session_state.compras)

        compras_aluno = compras[
            compras["Comprador"] == comprador_nome
        ]

        if len(compras_aluno) > 0:

            st.divider()
            st.subheader("📦 Compras desta sessão")

            st.dataframe(
                compras_aluno[
                    ["Código", "EDP", "Etapa", "Vendedor", "Preço"]
                ],
                use_container_width=True,
                hide_index=True
            )

            total = compras_aluno["Preço"].sum()

            st.metric("Total comprometido", f"{total:.0f} EC")

            st.warning(
                "ℹ️ Nesta etapa, as compras ainda ficam somente na sessão do "
                "Streamlit. O próximo módulo será a gravação permanente das "
                "compras e o desconto automático do saldo."
            )
